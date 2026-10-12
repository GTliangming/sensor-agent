"""文档解析：把多格式文件抽成「检索单元」列表（供 RAG 分块/向量化用）。

关键：表格类（Excel / Word 表格）按**行**切，每行自带表头，形成「型号: X；公开价: Y…」这样
自描述的记录，保证按型号精确检索。普通正文（PDF/Word/TXT/MD）则按自然文本返回、交给上层再切。

注意：产出的是“给机器检索的纯文本”，和前端“按原排版预览原始文件”是两条独立的路。
第一版支持 pdf / docx / xlsx / txt / md；老格式 .doc / .xls 不支持。
"""

from pathlib import Path

import docx
import openpyxl
import pypdf

SUPPORTED_EXTS = {"pdf", "docx", "xlsx", "txt", "md"}

# 这些格式产出的单元是「原子记录」（如表格一行），上层不要再按字符切碎
ATOMIC_EXTS = {"xlsx"}


def is_supported(ext: str) -> bool:
    return ext.lower().lstrip(".") in SUPPORTED_EXTS


def parse_units(path, ext: str) -> list[str]:
    """解析成检索单元列表；不支持的格式抛 ValueError。"""
    ext = ext.lower().lstrip(".")
    if ext == "pdf":
        return _pdf_units(path)
    if ext == "docx":
        return _docx_units(path)
    if ext == "xlsx":
        return _xlsx_units(path)
    if ext in ("txt", "md"):
        return [_read_text(path)]
    raise ValueError(f"暂不支持的格式: .{ext}（第一版支持 pdf/docx/xlsx/txt/md）")


def _pdf_units(path) -> list[str]:
    try:
        reader = pypdf.PdfReader(str(path))
    except Exception:
        raise ValueError("无法读取为 PDF（文件可能损坏或不是真正的 PDF）")
    # 一页一个单元，页内文本再由上层按长度切
    return [p.extract_text() or "" for p in reader.pages]


def _docx_units(path) -> list[str]:
    try:
        doc = docx.Document(str(path))
    except Exception:
        raise ValueError("无法读取为 .docx（可能是旧版 .doc 另存成 .docx，或文件损坏）")
    units: list[str] = []
    body = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if body.strip():
        units.append(body)
    # 表格：按行转成记录（docx 表格按单行表头处理，不传合并信息）
    for table in doc.tables:
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        units.extend(_rows_to_records(rows, []))
    return units


def _xlsx_units(path) -> list[str]:
    try:
        # 不用 read_only：否则拿不到合并单元格信息
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception:
        raise ValueError("无法读取为 .xlsx（可能是旧版 .xls 另存成 .xlsx，或文件损坏）")
    units: list[str] = []
    for ws in wb.worksheets:
        # 合并信息要在 unmerge 之前取（1-based 行列）
        merged = [
            (r.min_row, r.min_col, r.max_row, r.max_col)
            for r in ws.merged_cells.ranges
        ]
        _fill_merged(ws)  # 把合并单元格的值铺满整个合并区域
        rows = [
            [("" if c is None else str(c)).strip() for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
        units.extend(_rows_to_records(rows, merged, sheet=ws.title))
    wb.close()
    return units


def _fill_merged(ws) -> None:
    """把合并单元格左上角的值填充到整个合并区域，解决被合并格读出 None 的问题。"""
    for rng in list(ws.merged_cells.ranges):  # 先转 list，unmerge 会改动集合
        top_left = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                ws.cell(r, c).value = top_left


def _rows_to_records(rows, merged, sheet: str | None = None) -> list[str]:
    """多区块表格 → 记录列表，支持：分类横幅行、两级子表头、合并单元格。

    merged：unmerge 前的合并区域 [(min_row,min_col,max_row,max_col)]（1-based）。
    结构假设可反复出现：[分类横幅行] → [表头(1~2 行)] → [数据行...]。
    """
    records: list[str] = []
    section: str | None = None
    header: list[str] | None = None
    i, n = 0, len(rows)
    while i < n:
        vals = rows[i]
        nonempty = [v for v in vals if v]
        if not nonempty:
            i += 1
            continue
        # 分类横幅行（整行同值、跨 3 列以上）→ 开启新区块、重置表头
        if len(set(nonempty)) == 1 and len(nonempty) >= 3:
            section = nonempty[0]
            header = None
            i += 1
            continue
        if header is None:
            # 本行若有横向合并，说明下一行是子表头 → 表头占 2 行
            r0 = i + 1  # 1-based
            h = 2 if _has_horizontal_merge(merged, r0) else 1
            header = _compose_header(rows[i : i + h])
            i += h
            continue
        rec = _row_record(vals, header, section, sheet)
        if rec:
            records.append(rec)
        i += 1
    return records


def _has_horizontal_merge(merged, row: int) -> bool:
    return any(mr == row and xc > mc for (mr, mc, xr, xc) in merged)


def _compose_header(header_rows: list[list[str]]) -> list[str]:
    """多行表头合成复合列名：父·子（纵向合并父子同值时只保留一层）。

    无表头文字的列返回 ""（如最左侧的「类别」列往往没有列名）。
    """
    width = max((len(r) for r in header_rows), default=0)
    cols: list[str] = []
    for c in range(width):
        parts: list[str] = []
        for r in header_rows:
            v = r[c] if c < len(r) else ""
            if v and (not parts or parts[-1] != v):
                parts.append(v)
        cols.append("·".join(parts))  # 没有表头就是空串
    return cols


def _row_record(vals, header, section, sheet) -> str | None:
    pairs = []
    for idx, v in enumerate(vals):
        if not v:
            continue
        col = header[idx] if idx < len(header) else ""
        # 有列名用「列名: 值」，没列名（如最左类别列）直接放值，避免「列1:」这种无意义标签误导
        pairs.append(f"{col}: {v}" if col else v)
    if not pairs:
        return None
    prefix = f"[{sheet}] " if sheet else ""
    if section:
        prefix += f"分类: {section}；"
    return prefix + "；".join(pairs)


def _read_text(path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")
